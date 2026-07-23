import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(
    os.environ.get(
        "TEST_REPORT_OUT",
        Path(__file__).resolve().parents[1].parent / "SolarClean_pre_ship_test_report.docx",
    )
)

FONT = "Microsoft YaHei"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 90, 90)
LIGHT_BLUE = "E8EEF5"
WARN = "FFF2CC"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_para(par, before=0, after=6, line=1.15):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    set_para(p, after=0, line=1.1)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_dxa = int(sum(widths) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is not None:
        table._tbl.remove(tbl_grid)
    tbl_grid = OxmlElement("w:tblGrid")
    table._tbl.insert(1, tbl_grid)
    for width_in in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width_in * 1440)))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, width_in in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width_in)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width_in * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(table):
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for margin, width in (("top", "90"), ("bottom", "90"), ("start", "180"), ("end", "180")):
                node = tc_mar.find(qn(f"w:{margin}"))
                if node is None:
                    node = OxmlElement(f"w:{margin}")
                    tc_mar.append(node)
                node.set(qn("w:w"), width)
                node.set(qn("w:type"), "dxa")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_para(p, before=12 if level == 1 else 8, after=5)
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, 15, True, BLUE)
    elif level == 2:
        set_run_font(r, 12.5, True, BLUE)
    else:
        set_run_font(r, 11.5, True, DARK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_para(p, after=3, line=1.15)
        r = p.add_run(item)
        set_run_font(r, 10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_para(p, after=3, line=1.15)
        r = p.add_run(item)
        set_run_font(r, 10.5)


def add_callout(doc, title, text, fill=WARN):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_widths(table, [6.5])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    set_para(p, after=2, line=1.15)
    r = p.add_run(title + "：")
    set_run_font(r, 10.5, True, DARK)
    r = p.add_run(text)
    set_run_font(r, 10.5)
    set_cell_margins(table)


def add_check_table(doc, title, rows, widths=None):
    add_heading(doc, title, 2)
    widths = widths or [0.5, 1.55, 2.6, 0.9, 0.65]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    set_table_widths(table, widths)
    headers = ["序号", "测试项", "操作与判定", "结果", "备注"]
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], header, True, DARK, 9.5, WD_ALIGN_PARAGRAPH.CENTER)
    for idx, row in enumerate(rows, 1):
        cells = table.add_row().cells
        values = [str(idx), row[0], row[1], "□ 通过\n□ 失败", ""]
        for j, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[j], value, size=8.6, align=align)
    set_cell_margins(table)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Normal"].font.size = Pt(11)

    p = doc.add_paragraph()
    set_para(p, before=0, after=4, line=1.05)
    r = p.add_run("SolarClean 新增功能发货前测试报告")
    set_run_font(r, 20, True, RGBColor(11, 37, 69))

    p = doc.add_paragraph()
    set_para(p, after=8)
    r = p.add_run("覆盖：4G、上位机参数开关、遥控器在线/离线、断线继续喷水摆动；不覆盖固件升级/Bootloader。")
    set_run_font(r, 10.5, False, GRAY)

    p = doc.add_paragraph()
    set_para(p, after=10)
    r = p.add_run("建议测试时长：快速必测 60-90 分钟；有时间再做扩展项 30-60 分钟。")
    set_run_font(r, 10.5, True, DARK)

    add_callout(
        doc,
        "发货红线",
        "只要出现上电自动喷水/自动摆动、遥控器离线导致当前动作被改变、4G 无法上线、上位机参数掉电不保存、APP/上位机命令无响应或状态不可确认，建议不要发货。",
    )

    add_heading(doc, "1. 测试范围", 1)
    add_bullets(
        doc,
        [
            "必须测试：上电安全、上位机查询/保存 4G/遥控器/PSDK 开关、遥控器在线控制、遥控器关机/断连后保持当前动作、4G 在线、APP/云端控制水泵和摆动、可观察状态确认。",
            "不测试：固件升级、Bootloader 安装、OTA 回滚、长时间网络弱覆盖耐久。固件升级单独安排，不和本次发货前功能验收混在一起。",
            "测试目标：不是穷尽所有边界，而是在最短时间内把最容易导致现场事故的问题筛掉。",
        ],
    )

    add_heading(doc, "2. 测试环境准备", 1)
    prep_rows = [
        ("设备与安全", "水泵出口接安全回水/空载保护，确认喷头不会对人或设备喷射；舵机摆动区域无遮挡。"),
        ("遥控器", "确认遥控器、电池、接收机在线；记录 CH5 水泵、CH6 摆动、旋钮速度/幅度对应关系。"),
        ("4G", "插入 SIM 卡，天线连接可靠；确认当地信号正常，APP/云端能看到设备在线。"),
        ("上位机", "使用最新通用上位机，连接 USB 虚拟串口；打开高级设置弹窗，能查询/保存参数。"),
        ("测试记录", "测试人员记录上位机界面、APP/云端状态和实际动作；调试串口只作为研发复核的可选手段。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_widths(table, [1.45, 3.95, 0.8])
    for i, header in enumerate(["准备项", "要求", "确认"]):
        shade_cell(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], header, True, DARK, 9.5, WD_ALIGN_PARAGRAPH.CENTER)
    for name, req in prep_rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], name, True, size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], req, size=8.8)
        set_cell_text(cells[2], "□ OK", size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_margins(table)

    add_heading(doc, "3. 最短发货测试顺序", 1)
    add_numbered(
        doc,
        [
            "断开水泵危险输出或接回水，先做上电安全测试，确认不会自己启动。",
            "连接上位机，读取并保存 4G/遥控器/PSDK 开关参数，做一次断电重启验证。",
            "遥控器在线：手动控制水泵、摆动、速度/幅度，确认实际动作一致。",
            "遥控器离线：在水泵和摆动已启动状态下关闭遥控器，观察当前动作是否保持；再打开遥控器，确认不会因第一帧状态自己触发动作，必须拨动开关才重新控制。",
            "4G 开启：重启设备，看 APP/云端设备在线；下发水泵、摆动、速度/幅度命令，确认动作和状态反馈。",
            "4G 关闭：上位机关闭 4G 并重启，确认 APP/云端不显示设备通过 4G 上线，远程 4G 指令不应生效。",
            "PSDK 开关：开启/关闭后重启，确认 PSDK 启动或跳过符合参数，且不影响遥控器和 4G。",
        ],
    )

    add_heading(doc, "4. 必测项目记录", 1)
    add_check_table(
        doc,
        "4.1 上电安全与默认状态",
        [
            ("遥控器开关保持 ON 时上电", "设备上电后不得自动喷水或自动摆动；需要重新拨动遥控器开关后才允许动作。"),
            ("遥控器关闭时上电", "设备不得因为遥控器无信号而自动启动或异常停启；观察 30 秒无异常动作。"),
            ("水泵安全", "上电 30 秒内水泵继电器/压力 PWM 不应误动作。"),
        ],
    )
    add_check_table(
        doc,
        "4.2 上位机参数开关",
        [
            ("连接查询", "上位机连接成功后自动查询参数；界面显示 4G、遥控器、PSDK 开关当前值。"),
            ("保存到 Flash", "修改三个开关后保存，断电重启并重新连接上位机，界面显示值保持一致。"),
            ("遥控器开关", "关闭遥控器开关后，遥控器在线也不能控制水泵/摆动；重新开启后恢复。"),
            ("4G 开关", "关闭 4G 后重启，上位机参数保持关闭，APP/云端不应显示设备通过 4G 上线；开启后可恢复联网和远程控制。"),
            ("PSDK 开关", "关闭 PSDK 后重启，应看到 skip start_task 或无 PSDK 业务启动；开启后恢复。"),
        ],
    )
    add_check_table(
        doc,
        "4.3 遥控器在线控制",
        [
            ("CH5 水泵开关", "拨到 ON 水泵启动，拨到 OFF 水泵关闭；以继电器/水流/压力输出为准。"),
            ("CH6 摆动开关", "拨到 ON 摆动启动，拨到 OFF 摆动停止；以舵机实际动作确认。"),
            ("速度/幅度旋钮", "调节后摆动速度/幅度有明显变化；记录最小、中间、最大三个位置。"),
            ("反复拨动", "连续开关 5 次无卡死、无延迟失控、无异常复位。"),
        ],
    )
    add_check_table(
        doc,
        "4.4 遥控器离线继续喷水摆动",
        [
            ("先启动动作", "遥控器在线时开启水泵和摆动，确认动作稳定。"),
            ("遥控器主动关机", "关闭遥控器，观察 2 分钟：水泵和摆动保持当前状态，不允许被 CH5/CH6 低电平或全零帧关掉。"),
            ("可观察判定", "遥控器关闭后设备不应复位、不应停止当前动作、不应出现反复启停；如有指示灯/APP状态，以稳定为准。"),
            ("遥控器重新上线", "重新开遥控器后不应立刻触发开/关动作；必须实际拨动开关后才控制。"),
            ("上位机关闭遥控器控制", "如果上位机把“遥控器开关”关掉，则必须停止接受遥控器控制；这是参数开关，不是遥控器离线。"),
        ],
    )
    add_check_table(
        doc,
        "4.5 4G 与 APP/云端控制",
        [
            ("4G 启动流程", "开启 4G 后重启，APP/云端在 2 分钟内显示设备在线；若网络差可重试 1 次。"),
            ("水泵控制", "APP/云端下发水泵 ON/OFF，实际动作正确，APP/云端状态或回执正常。"),
            ("摆动控制", "APP/云端下发摆动 ON/OFF、速度、幅度，实际动作正确，APP/云端状态或回执正常。"),
            ("遥控器离线下 APP 控制", "遥控器关机时，APP/云端仍可控制水泵/摆动；遥控器离线帧不应覆盖 APP 当前控制。"),
            ("异常重连", "短暂断网或重启模块后，允许重连；设备不应复位或动作异常。"),
        ],
    )
    add_check_table(
        doc,
        "4.6 PSDK 与状态确认",
        [
            ("PSDK 开启", "飞机/PSDK 环境正常时组件可启动；无飞机时 PSDK 连接失败日志不能影响遥控器/4G。"),
            ("PSDK 关闭", "关闭后重启不启动 PSDK 任务；遥控器和 4G 功能仍正常。"),
            ("控制源隔离", "PSDK 开关变化不应影响遥控器和 4G 基本控制；用实际动作和 APP/上位机状态确认。"),
            ("无明显异常", "设备不应反复重启、反复上下线、指示状态异常闪烁或动作乱跳。"),
        ],
    )

    add_heading(doc, "5. 扩展项：有时间再测", 1)
    add_check_table(
        doc,
        "5.1 扩展稳定性",
        [
            ("连续运行", "4G 在线、水泵关闭、摆动关闭状态运行 30 分钟，无复位、无内存异常、无日志刷屏。"),
            ("连续喷水摆动", "水泵和摆动连续运行 10 分钟，观察电流、温升、自动保护是否符合预期。"),
            ("弱信号恢复", "人为制造弱信号/断网 1 次，恢复后 MQTT 能重新上线，动作状态不被乱改。"),
            ("多控制源切换", "遥控器控制后切 APP 控制，再关遥控器，确认最终动作由最后有效控制命令决定。"),
        ],
    )

    add_heading(doc, "6. 发货判定", 1)
    result_table = doc.add_table(rows=1, cols=3)
    result_table.style = "Table Grid"
    set_table_widths(result_table, [1.2, 4.1, 0.9])
    for i, header in enumerate(["结论", "条件", "签字"]):
        shade_cell(result_table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(result_table.rows[0].cells[i], header, True, DARK, 9.5, WD_ALIGN_PARAGRAPH.CENTER)
    rows = [
        ("可以发货", "4.1-4.6 必测全部通过；无红线问题；日志已保存。"),
        ("暂缓发货", "任意必测失败，或出现上电自启动、离线改变动作、4G 无法上线、参数不保存。"),
        ("带风险发货", "仅扩展项失败且客户场景可接受；必须记录风险和规避方式。"),
    ]
    for conclusion, condition in rows:
        cells = result_table.add_row().cells
        set_cell_text(cells[0], conclusion, True, size=9.2)
        set_cell_text(cells[1], condition, size=9.0)
        set_cell_text(cells[2], "", size=9.0)
    set_cell_margins(result_table)

    add_heading(doc, "7. 异常记录", 1)
    abnormal = doc.add_table(rows=1, cols=5)
    abnormal.style = "Table Grid"
    set_table_widths(abnormal, [0.5, 1.05, 2.35, 1.35, 0.95])
    for i, header in enumerate(["序号", "时间", "现象/日志", "处理结果", "负责人"]):
        shade_cell(abnormal.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(abnormal.rows[0].cells[i], header, True, DARK, 9.5, WD_ALIGN_PARAGRAPH.CENTER)
    for i in range(6):
        cells = abnormal.add_row().cells
        for j in range(5):
            set_cell_text(cells[j], str(i + 1) if j == 0 else "", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_margins(abnormal)

    add_heading(doc, "8. 现场建议", 1)
    add_bullets(
        doc,
        [
            "明天发货前优先跑 4.1-4.6，全部通过再考虑扩展项。",
            "测试人员优先保留上位机截图、APP/云端截图、异常视频或照片；调试串口日志由研发需要时再接。",
            "如果发现版本升级问题，先不要把 OTA 和功能测试混在一起判断；本报告只判定新增功能是否可发货。",
            "如果出现动作异常，先记录当时使用的是遥控器、APP/4G、上位机参数还是 PSDK，不要只写“异常”。",
        ],
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("SolarClean 发货前测试记录 - 仅覆盖新增功能，不覆盖固件升级")
    set_run_font(r, 8.5, False, GRAY)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
